"""
Headless Unit & Integration Test for Document Parsing & Batch Processing
"""

import os
from pathlib import Path
from src.document_parser import extract_text_from_file, split_into_chapters
from src.batch_processor import BatchProcessor
import docx

TEMP_DIR = Path("temp")
TEMP_DIR.mkdir(exist_ok=True)


def test_parsers():
    print("=== Testing Document Parsers ===")

    # 1. Text & Markdown
    sample_md = TEMP_DIR / "sample.md"
    sample_md.write_text(
        "# Einleitung\nDies ist der Einführungstext.\n\n# Kapitel 1: Die Reise\nHier beginnt das erste Kapitel.\n\n# Kapitel 2: Das Finale\nDas Ende der Geschichte.",
        encoding="utf-8"
    )
    extracted_md = extract_text_from_file(sample_md)
    print(f"[+] MD Extracted ({len(extracted_md)} chars):", extracted_md[:50], "...")

    chapters = split_into_chapters(extracted_md, "Buch")
    print(f"[+] MD Split into {len(chapters)} chapters:")
    for title, content in chapters:
        print(f"    - [{title}] ({len(content)} chars)")
    assert len(chapters) == 3, f"Expected 3 chapters, got {len(chapters)}"

    # 2. DOCX
    sample_docx = TEMP_DIR / "sample.docx"
    doc = docx.Document()
    doc.add_heading("Kapitel 1: Grundlagen", level=1)
    doc.add_paragraph("Dies ist ein Absatz in Microsoft Word.")
    doc.add_heading("Kapitel 2: Praxis", level=1)
    doc.add_paragraph("Praktische Anwendung mit Gemini TTS.")
    doc.save(str(sample_docx))

    extracted_docx = extract_text_from_file(sample_docx)
    print(f"\n[+] DOCX Extracted ({len(extracted_docx)} chars):", extracted_docx[:50], "...")
    docx_chapters = split_into_chapters(extracted_docx, "WordDoc")
    print(f"[+] DOCX Split into {len(docx_chapters)} chapters:")
    for title, content in docx_chapters:
        print(f"    - [{title}] ({len(content)} chars)")
    assert len(docx_chapters) == 2, f"Expected 2 chapters, got {len(docx_chapters)}"

    # 3. SRT Subtitles
    sample_srt = TEMP_DIR / "sample.srt"
    sample_srt.write_text(
        "1\n00:00:01,000 --> 00:00:03,500\nHallo und herzlich willkommen!\n\n2\n00:00:04,000 --> 00:00:07,200\nHeute testen wir den Batch Import.\n",
        encoding="utf-8"
    )
    extracted_srt = extract_text_from_file(sample_srt)
    print(f"\n[+] SRT Extracted ({len(extracted_srt)} chars):", extracted_srt)
    assert "Hallo und herzlich willkommen" in extracted_srt
    assert "00:00:" not in extracted_srt

    # 4. BatchProcessor Queue
    bp = BatchProcessor()
    items = bp.add_file(sample_md, split_chapters=True)
    print(f"\n[+] BatchProcessor added {len(items)} items from MD file.")
    assert len(bp.items) == 3

    print("\n[+] ALL DOCUMENT PARSER & CHAPTER TESTS PASSED SUCCESSFULLY!")


if __name__ == "__main__":
    test_parsers()
